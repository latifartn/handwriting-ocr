"""
Handwritten PDF -> structured Excel

Pipeline (engine=vlm):
    PDF page -> image -> VLM transcription (image -> text only)
             -> text-only LLM extraction (with per-value evidence)
             -> evidence validation (anti-hallucination)
             -> Excel (Data / Evidence / Raw OCR sheets)

Requirements:
    pip install pymupdf opencv-python numpy pillow openpyxl
    (TrOCR engine only: pip install transformers torch)

Ollama must be running and BOTH models must be pulled, e.g.:
    ollama pull qwen3-vl:8b
    ollama pull qwen2.5:7b
"""

import argparse
import base64
import io
import json
import re
import urllib.error
import urllib.request
from pathlib import Path

import cv2
import numpy as np
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from PIL import Image

try:
    import pymupdf  # new name
except ImportError:  # old PyMuPDF
    import fitz as pymupdf


# ============================================================
# PDF -> images
# ============================================================

def pdf_pages(path, dpi=220):
    doc = pymupdf.open(str(path))
    scale = dpi / 72.0
    mat = pymupdf.Matrix(scale, scale)

    pages = []
    for page in doc:
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img = np.frombuffer(pix.samples, dtype=np.uint8)

        if pix.n == 1:
            img = img.reshape(pix.height, pix.width)
            img = cv2.cvtColor(img, cv2.COLOR_GRAY2RGB)
        elif pix.n == 3:
            img = img.reshape(pix.height, pix.width, 3)
        elif pix.n == 4:
            img = img.reshape(pix.height, pix.width, 4)
            img = cv2.cvtColor(img, cv2.COLOR_RGBA2RGB)
        else:
            raise ValueError(f"Unsupported pixmap channel count: {pix.n}")

        pages.append(np.ascontiguousarray(img))

    doc.close()
    return pages


# ============================================================
# TrOCR (optional engine)
# ============================================================

class TrOCR:
    def __init__(self, model_name="microsoft/trocr-base-handwritten"):
        import torch
        from transformers import TrOCRProcessor, VisionEncoderDecoderModel

        self.torch = torch
        self.processor = TrOCRProcessor.from_pretrained(model_name)
        self.model = VisionEncoderDecoderModel.from_pretrained(model_name)
        self.device = "cuda" if torch.cuda.is_available() else "cpu"
        self.model.to(self.device)
        self.model.eval()

    def read(self, image):
        if isinstance(image, np.ndarray):
            image = Image.fromarray(image)

        pixel_values = self.processor(
            images=image, return_tensors="pt"
        ).pixel_values.to(self.device)

        with self.torch.no_grad():
            ids = self.model.generate(pixel_values, max_new_tokens=128)

        text = self.processor.batch_decode(ids, skip_special_tokens=True)[0]
        return text.strip(), 1.0


# ============================================================
# Image utils
# ============================================================

def is_blank(img, threshold=245):
    gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
    return float(np.mean(gray > threshold)) > 0.995


def pad_white(img, pad=10):
    return cv2.copyMakeBorder(
        img, pad, pad, pad, pad, cv2.BORDER_CONSTANT, value=(255, 255, 255)
    )


def clean_page(page):
    gray = cv2.cvtColor(page, cv2.COLOR_RGB2GRAY)
    gray = cv2.GaussianBlur(gray, (3, 3), 0)  # mild denoise only
    return cv2.cvtColor(gray, cv2.COLOR_GRAY2RGB)


def encode_jpeg(image, max_side=1800):
    h, w = image.shape[:2]
    scale = min(1.0, max_side / max(h, w))

    if scale < 1.0:
        image = cv2.resize(
            image, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA
        )

    buf = io.BytesIO()
    Image.fromarray(image).save(buf, format="JPEG", quality=92)
    return base64.b64encode(buf.getvalue()).decode("utf-8")


# ============================================================
# Table cell detection (TrOCR cells mode)
# ============================================================

def detect_cells(page):
    gray = cv2.cvtColor(page, cv2.COLOR_RGB2GRAY)
    bw = cv2.threshold(gray, 200, 255, cv2.THRESH_BINARY_INV)[1]

    horizontal = cv2.morphologyEx(
        bw, cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (max(20, page.shape[1] // 20), 1)),
    )
    vertical = cv2.morphologyEx(
        bw, cv2.MORPH_OPEN,
        cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(20, page.shape[0] // 20))),
    )

    lines = cv2.bitwise_or(horizontal, vertical)
    contours, _ = cv2.findContours(lines, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)

    boxes = []
    for c in contours:
        x, y, w, h = cv2.boundingRect(c)
        if w < 30 or h < 20 or w * h < 1500:
            continue
        boxes.append((x, y, w, h))
    return boxes


def boxes_to_grid(boxes):
    if not boxes:
        return []

    boxes = sorted(boxes, key=lambda b: (b[1], b[0]))
    rows = []

    for box in boxes:
        _, y, _, h = box
        for row in rows:
            ry = np.mean([b[1] for b in row])
            if abs(y - ry) < max(20, h * 0.5):
                row.append(box)
                break
        else:
            rows.append([box])

    return [sorted(row, key=lambda b: b[0]) for row in rows]


def run_cells(page, trocr, debug_dir=None):
    rows = boxes_to_grid(detect_cells(page))
    output = []
    crops_dbg = []

    for row in rows:
        row_values = []
        for x, y, w, h in row:
            crop = page[max(0, y):min(page.shape[0], y + h),
                        max(0, x):min(page.shape[1], x + w)]
            crop = pad_white(crop)
            crops_dbg.append(crop)

            if is_blank(crop):
                row_values.append(("", 0.0))
                continue

            row_values.append(trocr.read(crop))
        output.append(row_values)

    if debug_dir:
        dump_crops(crops_dbg, debug_dir)

    return output


def dump_crops(crops, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for i, crop in enumerate(crops):
        cv2.imwrite(str(out_dir / f"crop_{i:04d}.png"),
                    cv2.cvtColor(crop, cv2.COLOR_RGB2BGR))


# ============================================================
# Line splitting (TrOCR lines mode)
# ============================================================

def split_lines(page):
    gray = cv2.cvtColor(page, cv2.COLOR_RGB2GRAY)
    ink = 255 - gray
    projection = np.sum(ink > 40, axis=1)
    active = projection > max(3, int(page.shape[1] * 0.005))

    spans = []
    start = None
    for i, flag in enumerate(active):
        if flag and start is None:
            start = i
        elif not flag and start is not None:
            if i - start >= 8:
                spans.append((start, i))
            start = None
    if start is not None and len(active) - start >= 8:
        spans.append((start, len(active)))

    crops = []
    for y1, y2 in spans:
        crop = page[max(0, y1 - 5):min(page.shape[0], y2 + 5), :]
        if not is_blank(crop):
            crops.append(crop)
    return crops


def run_lines(page, trocr):
    lines = []
    for crop in split_lines(page):
        text, conf = trocr.read(crop)
        if text.strip():
            lines.append({"text": text.strip(), "confidence": conf})
    return lines


# ============================================================
# Ollama HTTP
# ============================================================

NUM_CTX = 8192  # Ollama default (2-4k) silently truncates long transcripts


def ollama_chat(messages, model, url, temperature=0, fmt="json",
                timeout=900, think=False, extra_options=None):
    payload = {
        "model": model,
        "messages": messages,
        "stream": False,
        "think": think,  # ignored by non-thinking models / older servers
        "options": {"temperature": temperature, "num_ctx": NUM_CTX},
    }
    if extra_options:
        payload["options"].update(extra_options)
    if fmt:
        payload["format"] = fmt

    req = urllib.request.Request(
        url.rstrip("/") + "/api/chat",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            result = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        body = ""
        try:
            body = e.read().decode("utf-8", errors="replace")
        except Exception:
            pass
        hint = ""
        if e.code == 404:
            hint = (f"\n  -> Model '{model}' is probably not pulled. "
                    f"Run:  ollama pull {model}")
        raise RuntimeError(f"Ollama HTTP {e.code} for model '{model}': {body}{hint}") from None
    except urllib.error.URLError as e:
        raise RuntimeError(
            f"Cannot reach Ollama at {url} ({e.reason}). Is `ollama serve` running?"
        ) from None

    return result.get("message", {}).get("content", "")


def list_ollama_models(url):
    req = urllib.request.Request(url.rstrip("/") + "/api/tags")
    with urllib.request.urlopen(req, timeout=15) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return [m.get("name", "") for m in data.get("models", [])]


def model_available(wanted, installed):
    if wanted in installed:
        return True
    if ":" not in wanted:
        return f"{wanted}:latest" in installed
    return False


def preflight(args):
    """Fail early with a clear message instead of a bare 404 per PDF."""
    needed = []
    if args.engine == "vlm":
        needed.append(args.vlm_model)
        if args.mode != "text":
            needed.append(args.ollama_model)
    elif not args.no_llm:
        needed.append(args.ollama_model)

    if not needed:
        return

    try:
        installed = list_ollama_models(args.ollama_url)
    except Exception as e:
        raise SystemExit(
            f"Can't talk to Ollama at {args.ollama_url}: {e}\n"
            f"Start it with `ollama serve` (or launch the Ollama app)."
        )

    missing = [m for m in needed if not model_available(m, installed)]
    if missing:
        cmds = "\n".join(f"    ollama pull {m}" for m in missing)
        raise SystemExit(
            "Missing Ollama model(s): " + ", ".join(missing) + "\n"
            "Pull them first:\n" + cmds + "\n\n"
            "Installed: " + (", ".join(installed) or "(none)")
        )


# ============================================================
# Evidence validation
# ============================================================

def normalize_match(text):
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text).casefold()).strip()


def strip_line_prefix(text):
    """Remove the '12: ' numbering we add to OCR lines, if the model copied it."""
    return re.sub(r"^\s*\d+\s*:\s*", "", str(text or ""))


def evidence_supported(value, evidence, transcript):
    value_n = normalize_match(value)
    evidence_n = normalize_match(strip_line_prefix(evidence))
    transcript_n = normalize_match(transcript)

    if not value_n or not evidence_n:
        return False
    if evidence_n not in transcript_n:
        return False
    return value_n in evidence_n


# ============================================================
# Text-only extraction
# ============================================================

EXTRACT_SYSTEM = """
You are a strict OCR data extraction system.
Your ONLY source of truth is the OCR text provided by the user.

ABSOLUTE RULES:
1. NEVER guess, infer, or use outside knowledge.
2. NEVER complete missing information or derive one field from another.
3. NEVER assume a location, country, state, city, ZIP code, date, name, etc.
4. If a value is not explicitly in the OCR text, or is ambiguous/illegible, return "".
5. Evidence MUST be copied exactly from a single OCR line and MUST contain the value.
6. Do not invent evidence.
7. Preserve the value as written in the OCR.
8. Do not turn contextual info into a field value.

Example:
OCR: "Essex"            -> Country = ""   (USA was never written)
OCR: "State: Essex"     -> State = "Essex", Evidence = "State: Essex"

Return JSON only.
"""


def number_lines(transcript):
    out = []
    n = 1
    for line in transcript.split("\n"):
        line = line.strip()
        if not line:
            continue
        out.append(f"{n}: {line}")
        n += 1
    return "\n".join(out)


def ollama_extract_from_text(transcript, columns, model, url):
    if not transcript.strip():
        return [], "EMPTY OCR TRANSCRIPT"

    numbered_text = number_lines(transcript)
    cols = [c.strip() for c in (columns or []) if c.strip()]
    field_value_mode = (not cols) or cols == ["Field", "Value"]

    if field_value_mode:
        instruction = """
Identify each label/value pair that is explicitly written in the OCR text.
Return one row per pair.

Example OCR:
Name: John
City: Riga

Example output:
{"rows": [
  {"Field": "Name", "Value": "John", "Evidence": "Name: John"},
  {"Field": "City", "Value": "Riga", "Evidence": "City: Riga"}
]}

Do NOT add fields that are not written.
"""
    else:
        example = {c: "..." for c in cols}
        for c in cols:
            example[f"{c}__evidence"] = "exact OCR line"
        instruction = f"""
Extract ONLY these columns: {cols}

Return ONE row. For each column give the value ("" if not explicitly present)
and, under "<column>__evidence", the exact OCR line containing that value
("" if the value is "").

Format:
{json.dumps({"rows": [example]}, ensure_ascii=False, indent=2)}
"""

    user = f"{instruction}\n\nOCR TEXT:\n\n{numbered_text}\n"

    msgs = [{"role": "system", "content": EXTRACT_SYSTEM},
            {"role": "user", "content": user}]
    attempts = [
        dict(temperature=0, extra_options=None),
        dict(temperature=0.2, extra_options={"repeat_penalty": 1.15, "num_predict": 2048}),
        dict(temperature=0.4, extra_options={"repeat_penalty": 1.3, "num_predict": 1024}),
    ]
    raw, last_err = None, None
    for att in attempts:
        try:
            raw = ollama_chat(msgs, model=model, url=url, fmt="json", **att)
            break
        except RuntimeError as e:
            last_err = e
            print(f"    extraction retry ({str(e)[:80]}...)")
    if raw is None:
        return [], f"EXTRACTION FAILED: {last_err}"

    try:
        data = json.loads(raw)
    except Exception:
        return [], raw

    rows = data.get("rows", []) if isinstance(data, dict) else []
    if not isinstance(rows, list):
        return [], raw

    validated = []
    for row in rows:
        if not isinstance(row, dict):
            continue

        if field_value_mode:
            field = str(row.get("Field", "")).strip()
            value = str(row.get("Value", "")).strip()
            evidence = str(row.get("Evidence", "")).strip()

            if value and not evidence_supported(value, evidence, transcript):
                value, evidence = "", ""

            if not field and not value:
                continue

            validated.append({"Field": field, "Value": value, "_evidence": evidence})
        else:
            out = {}
            ev = {}
            for c in cols:
                value = str(row.get(c, "") or "").strip()
                evidence = str(row.get(f"{c}__evidence", "") or "").strip()
                if value and not evidence_supported(value, evidence, transcript):
                    value, evidence = "", ""
                out[c] = value
                ev[c] = evidence
            out["_evidence"] = json.dumps(ev, ensure_ascii=False)
            validated.append(out)

    return validated, raw


# Legacy path for TrOCR lines mode
def ollama_rows(lines, columns, model, url):
    cols = [c.strip() for c in columns if c.strip()]
    numbered = "\n".join(f"{i}: {line}" for i, line in enumerate(lines, 1))

    system = """
You convert OCR text into structured data.
Use ONLY information explicitly present in the OCR. Never guess, infer or use
outside knowledge. Missing/uncertain fields must be "". Preserve OCR values.
Return JSON only.
"""
    example = {c: "value" for c in cols} if cols else {"Field": "name", "Value": "value"}
    user = f"""Columns: {cols}

OCR:
{numbered}

Return: {json.dumps({"rows": [example]}, ensure_ascii=False)}
"""
    raw = ollama_chat(
        [{"role": "system", "content": system}, {"role": "user", "content": user}],
        model=model, url=url, temperature=0, fmt="json",
    )

    try:
        data = json.loads(raw)
    except Exception:
        return [], raw

    rows = data.get("rows", []) if isinstance(data, dict) else []
    if not isinstance(rows, list):
        return [], raw
    return [r for r in rows if isinstance(r, dict)], raw


# ============================================================
# VLM: image -> text ONLY
# ============================================================

TRANSCRIBE_SYSTEM = """
You are a strict OCR transcription system.
Your ONLY task is to transcribe text that is actually visible in the image.

Do NOT extract fields, interpret, summarize, classify, infer, or guess.
Transcribe line by line, preserving spelling, numbers, dates, punctuation,
names, addresses, labels, handwritten and printed text.
Use a newline between lines of text.
If a character or word is genuinely unreadable, use [?].
Never invent missing text.

Output ONLY the plain transcribed text. No JSON, no markdown, no commentary.
"""


def vlm_transcribe_pages(pages, args, debug_dir=None, stem="doc"):
    transcripts = []

    for page_number, page in enumerate(pages, start=1):
        cleaned = clean_page(page)

        if debug_dir:
            Path(debug_dir).mkdir(parents=True, exist_ok=True)
            cv2.imwrite(
                str(Path(debug_dir) / f"{stem}_p{page_number}.png"),
                cv2.cvtColor(cleaned, cv2.COLOR_RGB2BGR),
            )

        b64 = encode_jpeg(cleaned, max_side=args.vlm_max_side)

        # Ollama native format: images go in message["images"] as raw base64
        messages = [
            {"role": "system", "content": TRANSCRIBE_SYSTEM},
            {
                "role": "user",
                "content": "Transcribe this page exactly. Do not extract or infer fields.",
                "images": [b64],
            },
        ]

        raw = ollama_chat(
            messages, model=args.vlm_model, url=args.ollama_url,
            temperature=0, fmt=None, timeout=900,
        )

        transcript = raw.strip()
        # strip <think> blocks / JSON wrapper if the model added them anyway
        transcript = re.sub(r"<think>.*?</think>", "", transcript, flags=re.S).strip()
        if transcript.startswith("{"):
            try:
                t = json.loads(transcript).get("transcript", transcript)
                transcript = ("\n".join(map(str, t)) if isinstance(t, list) else str(t)).strip()
            except Exception:
                pass
        transcript = re.sub(r"^```\w*\n|\n```$", "", transcript).strip()

        if not transcript:
            print(f"    WARNING: VLM returned empty text for page {page_number} "
                  f"(model={args.vlm_model}). Try --vlm-model qwen2.5vl:7b")

        if debug_dir:
            (Path(debug_dir) / f"{stem}_p{page_number}.txt").write_text(
                transcript, encoding="utf-8"
            )

        transcripts.append(transcript)

    return transcripts


def run_vlm(pages, args, columns, stem="doc"):
    page_transcripts = vlm_transcribe_pages(pages, args, args.debug or None, stem)

    full_transcript = "\n\n".join(
        f"PAGE {i}\n{text}" for i, text in enumerate(page_transcripts, start=1)
    )

    rows, extraction_raw = ollama_extract_from_text(
        full_transcript, columns, args.ollama_model, args.ollama_url
    )

    # Fallback: free-form pages with no label/value pairs -> one transcript line per row
    fv_mode = (not columns) or [c.strip() for c in columns] == ["Field", "Value"]
    has_values = any(str(r.get("Value", "")).strip() for r in rows)
    if fv_mode and not has_values:
        fallback = []
        for page_i, text in enumerate(page_transcripts, start=1):
            for line in text.split("\n"):
                line = line.strip()
                if line:
                    fallback.append({"Field": f"Page {page_i}", "Value": line,
                                     "_evidence": line})
        if fallback:
            print("    note: no Field/Value pairs found, wrote transcript lines instead")
            rows = fallback

    raw = (
        "===== VLM TRANSCRIPTION =====\n\n" + full_transcript
        + "\n\n===== EXTRACTION RAW RESPONSE =====\n\n" + extraction_raw
        + "\n\n===== VALIDATED ROWS =====\n\n"
        + json.dumps(rows, ensure_ascii=False, indent=2)
    )
    return rows, raw


def run_text_mode(pages, args, stem="doc"):
    transcripts = vlm_transcribe_pages(pages, args, args.debug or None, stem)
    return "\n\n".join(
        f"===== PAGE {i} =====\n{t}" for i, t in enumerate(transcripts, start=1)
    )


# ============================================================
# Excel writer
# ============================================================

def clean_cell(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]


def set_text(ws, row, col, value):
    """Write value as a plain string. Never let Excel treat it as a formula."""
    cell = ws.cell(row=row, column=col)
    cell.value = clean_cell(value)
    cell.data_type = "s"
    return cell


def write_xlsx(path, rows, columns=None, raw_ocr=""):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    if columns:
        headers = list(columns)
    elif rows:
        headers = [k for k in rows[0].keys() if not k.startswith("_")]
    else:
        headers = ["Field", "Value"]

    for ci, h in enumerate(headers, start=1):
        set_text(ws, 1, ci, h)

    for ri, row in enumerate(rows, start=2):
        for ci, h in enumerate(headers, start=1):
            set_text(ws, ri, ci, row.get(h, ""))

    # Evidence sheet (only if any row carries evidence)
    if any(r.get("_evidence") for r in rows):
        ev = wb.create_sheet("Evidence")
        ev.cell(row=1, column=1, value="Row")
        ev.cell(row=1, column=2, value="Evidence")
        for i, r in enumerate(rows, start=1):
            ev.cell(row=i + 1, column=1, value=i)
            set_text(ev, i + 1, 2, r.get("_evidence", ""))

    raw_ws = wb.create_sheet("Raw OCR")
    for ri, line in enumerate(str(raw_ocr).split("\n"), start=1):
        set_text(raw_ws, ri, 1, line)
    raw_ws.column_dimensions["A"].width = 120

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.title == "Raw OCR":
            continue
        for col_cells in sheet.columns:
            longest = max((len(str(c.value)) for c in col_cells if c.value is not None),
                          default=0)
            sheet.column_dimensions[col_cells[0].column_letter].width = max(
                12, min(longest, 50) + 2
            )

    wb.save(path)


# ============================================================
# Process one PDF
# ============================================================

def process_pdf(pdf_path, args):
    pages = pdf_pages(pdf_path, dpi=args.dpi)
    stem = Path(pdf_path).stem

    columns = [x.strip() for x in args.columns.split(",") if x.strip()] \
        if args.columns else []

    # ---- TEXT MODE ----
    if args.mode == "text":
        if args.engine != "vlm":
            raise ValueError("--mode text requires --engine vlm")
        return [], run_text_mode(pages, args, stem), None

    # ---- VLM ----
    if args.engine == "vlm":
        rows, raw = run_vlm(pages, args, columns, stem)
        return rows, raw, (columns or ["Field", "Value"])

    # ---- TrOCR ----
    trocr = TrOCR(args.trocr_model)

    if args.mode == "lines":
        all_lines = []
        for page in pages:
            all_lines.extend(item["text"] for item in run_lines(page, trocr))

        if args.no_llm:
            rows = [{"Field": "Text", "Value": line} for line in all_lines]
            return rows, "\n".join(all_lines), ["Field", "Value"]

        rows, raw = ollama_rows(all_lines, columns, args.ollama_model, args.ollama_url)
        return rows, raw, (columns or None)

    # cells mode
    all_rows = []
    for page in pages:
        all_rows.extend(run_cells(page, trocr, args.debug or None))

    if not columns:
        max_cols = max((len(r) for r in all_rows), default=0)
        columns = [f"Column{i + 1}" for i in range(max_cols)]

    rows = []
    for row in all_rows:
        out = {}
        for i, value in enumerate(row):
            if i >= len(columns):
                break
            text = value[0] if isinstance(value, tuple) else str(value)
            out[columns[i]] = text
        rows.append(out)

    return rows, "", columns


# ============================================================
# Files / CLI
# ============================================================

def find_pdfs(input_path):
    path = Path(input_path)
    if path.is_file():
        return [path] if path.suffix.lower() == ".pdf" else []
    if path.is_dir():
        return sorted(path.glob("*.pdf"))
    return []


def main():
    p = argparse.ArgumentParser(description="Handwritten PDF OCR -> structured Excel")
    p.add_argument("input", help="PDF file or directory containing PDFs")
    p.add_argument("--engine", choices=["trocr", "vlm"], default="trocr")
    p.add_argument("--mode", choices=["cells", "lines", "text"], default="cells")
    p.add_argument("--columns", default="", help='Comma-separated, e.g. "Field,Value"')
    p.add_argument("--out-dir", default="out")
    p.add_argument("--debug", default="", help="Dir for page images / transcripts / crops")
    p.add_argument("--dpi", type=int, default=220)
    p.add_argument("--trocr-model", default="microsoft/trocr-base-handwritten")
    p.add_argument("--vlm-model", default="qwen3-vl:8b")
    p.add_argument("--vlm-max-side", type=int, default=1800)
    p.add_argument("--ollama-model", default="qwen2.5:7b")
    p.add_argument("--ollama-url", default="http://localhost:11434")
    p.add_argument("--no-llm", action="store_true")
    p.add_argument("--num-ctx", type=int, default=8192,
                   help="Ollama context window; raise for long multi-page scans")
    args = p.parse_args()
    global NUM_CTX
    NUM_CTX = args.num_ctx

    pdfs = find_pdfs(args.input)
    if not pdfs:
        print(f"No PDF files found: {args.input}")
        return

    preflight(args)

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    if args.debug:
        Path(args.debug).mkdir(parents=True, exist_ok=True)

    ok = failed = 0

    for i, pdf in enumerate(pdfs, start=1):
        print(f"[{i}/{len(pdfs)}] Processing {pdf.name}...")
        try:
            rows, raw, columns = process_pdf(pdf, args)

            if args.mode == "text":
                out = out_dir / f"{pdf.stem}.txt"
                out.write_text(raw, encoding="utf-8")
                print(f"    -> {out}")
            else:
                out = out_dir / f"{pdf.stem}.xlsx"
                write_xlsx(out, rows, columns=columns, raw_ocr=raw)
                print(f"    -> {out}")
                filled = sum(1 for r in rows
                             if any(str(v).strip() for k, v in r.items()
                                    if not k.startswith("_") and k != "Field"))
                print(f"    rows: {len(rows)} ({filled} with values)")
                if args.engine == "vlm":
                    m = re.search(r"===== VLM TRANSCRIPTION =====\n\n(.*?)\n\n===== EXTRACTION",
                                  raw, re.S)
                    body = re.sub(r"PAGE \d+", "", m.group(1)).strip() if m else ""
                    if not body:
                        print("    WARNING: VLM transcript is EMPTY -> check dbg page image/txt")
                    elif not filled:
                        print("    WARNING: transcript has text but 0 values survived "
                              "extraction/validation -> see 'Raw OCR' sheet")
            ok += 1

        except Exception as e:
            failed += 1
            print(f"    FAILED: {e}")

    print(f"\n[done] {ok} ok, {failed} failed")


if __name__ == "__main__":
    main()